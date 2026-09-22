import argparse
import json
import logging
import os
import re
import sys
import uuid
from datetime import datetime
from pathlib import Path

import openpyxl
import pyodbc
from dotenv import load_dotenv


def load_env_files() -> None:
    root_dir = Path(__file__).resolve().parents[1]
    load_dotenv(root_dir / ".env")
    load_dotenv(root_dir / "ETL" / ".env")


load_env_files()

SERVER = os.getenv("DB_SERVER")
DATABASE = os.getenv("DB_NAME")
USER = os.getenv("DB_USER")
PASSWORD = os.getenv("DB_PASSWORD")
DRIVER_ENV = os.getenv("DB_DRIVER")

TABLE = "dbo.tbl_gestiones_diarias_sct"
BATCH_SIZE = 1000
LOGS_DIR = Path(__file__).resolve().parent / "logs"
DEFAULT_INPUT_DIR = Path(r"C:\Users\Explotación Phoenix\Desktop\PHOENIX 2\PASO\Gestion SCT")
DEFAULT_FILE_RE = re.compile(r"^(\d{8}) - Gestiones del dia - PHOENIX\.(xlsx|xlsm)$", re.IGNORECASE)

EXPECTED_COLUMNS = [
    ("PERIODO", "periodo"),
    ("DDAS_FEC_PROC", "ddas_fec_proc"),
    ("DDAS_NRT_PPAL", "ddas_nrt_ppal"),
    ("DDAS_DRT_PPAL", "ddas_drt_ppal"),
    ("DDAS_ID_NUMERO_OPERAC", "ddas_id_numero_operac"),
    ("TRAMO_MORA", "tramo_mora"),
    ("COBRADOR_ACTUAL", "cobrador_actual"),
    ("EMPRESA", "empresa"),
    ("FECH_GEST", "fech_gest"),
    ("FECH_INGR", "fech_ingr"),
    ("CONTACTO", "contacto"),
    ("ESTADO", "estado"),
    ("FECHA_COMP", "fecha_comp"),
    ("COM_GEST", "com_gest"),
    ("TIPO_GEST", "tipo_gest"),
    ("CLASIFICACIONGESTION", "clasificacion_gestion"),
    ("ACCIÓN", "accion"),
    ("CANAL", "canal"),
    ("ZONA", "zona"),
]

DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S.%f",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%d-%m-%Y %H:%M:%S.%f",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
    "%d-%m-%Y",
    "%d/%m/%Y %H:%M:%S.%f",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y",
)


def normalize_header(value: object) -> str:
    text = str(value or "").replace("\ufeff", "").replace("\x00", "").strip().upper()
    text = text.replace("Á", "A").replace("É", "E").replace("Í", "I").replace("Ó", "O").replace("Ú", "U")
    return re.sub(r"[^A-Z0-9]+", "_", text).strip("_")


EXPECTED_HEADER_KEYS = {normalize_header(header): (header, sql_name) for header, sql_name in EXPECTED_COLUMNS}


def pick_driver() -> str:
    available = list(pyodbc.drivers())
    preferred = []
    if DRIVER_ENV:
        preferred.append(DRIVER_ENV)
    preferred += ["ODBC Driver 18 for SQL Server", "ODBC Driver 17 for SQL Server", "SQL Server"]
    for driver in preferred:
        if driver in available:
            return driver
    raise RuntimeError(f"No hay driver ODBC para SQL Server. Drivers encontrados: {available}")


def connect() -> pyodbc.Connection:
    missing = [
        name
        for name, value in {
            "DB_SERVER": SERVER,
            "DB_NAME": DATABASE,
            "DB_USER": USER,
            "DB_PASSWORD": PASSWORD,
        }.items()
        if not value
    ]
    if missing:
        raise RuntimeError("Faltan variables en .env: " + ", ".join(missing))

    driver = pick_driver()
    conn_str = (
        f"Driver={{{driver}}};"
        f"Server={SERVER};"
        f"Database={DATABASE};"
        f"Uid={USER};"
        f"Pwd={PASSWORD};"
        "TrustServerCertificate=yes;"
    )
    if driver != "SQL Server":
        conn_str += "Encrypt=yes;"
    return pyodbc.connect(conn_str)


def sql_ident(name: str) -> str:
    return "[" + str(name).replace("]", "]]") + "]"


def setup_logger() -> logging.Logger:
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("etl_gestiones_diarias_sct")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    log_file = LOGS_DIR / f"etl_gestiones_diarias_sct_{datetime.now().strftime('%Y%m%d')}.log"
    handler = logging.FileHandler(log_file, encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s | %(levelname)s | %(message)s"))
    logger.addHandler(handler)
    return logger


def is_empty_value(value: object) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text == "" or text in {"NaN", "nan", "NaT", "None", "NULL"}


def value_to_text(value: object) -> str | None:
    if is_empty_value(value):
        return None
    if isinstance(value, datetime):
        return value.replace(microsecond=0).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\x00", "")


def parse_fecha_gestion(value: object) -> datetime | None:
    if is_empty_value(value):
        return None
    if isinstance(value, datetime):
        return value

    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def validate_path(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo: {path}")
    if not path.is_file():
        raise RuntimeError(f"La ruta no corresponde a un archivo: {path}")
    if path.stat().st_size == 0:
        raise RuntimeError(f"El archivo esta vacio: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise RuntimeError(f"Extension no permitida: {path.suffix}")


def resolve_input_file(file_path: str | None) -> Path:
    if file_path:
        return Path(file_path)

    if not DEFAULT_INPUT_DIR.exists():
        raise FileNotFoundError(f"No existe la carpeta por defecto: {DEFAULT_INPUT_DIR}")
    if not DEFAULT_INPUT_DIR.is_dir():
        raise RuntimeError(f"La ruta por defecto no es una carpeta: {DEFAULT_INPUT_DIR}")

    candidates: list[tuple[str, Path]] = []
    for path in DEFAULT_INPUT_DIR.iterdir():
        if not path.is_file():
            continue
        match = DEFAULT_FILE_RE.match(path.name)
        if match:
            candidates.append((match.group(1), path))

    if not candidates:
        raise FileNotFoundError(
            "No se encontro archivo en "
            f"{DEFAULT_INPUT_DIR} con patron YYYYMMDD - Gestiones del dia - PHOENIX.xlsx"
        )

    return max(candidates, key=lambda item: item[0])[1]


def find_source_sheet(workbook) -> tuple[str, dict[str, int], int]:
    best_missing: list[str] | None = None
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True), start=1):
            headers = [normalize_header(value) for value in row]
            non_empty = [header for header in headers if header]
            duplicates = sorted({header for header in non_empty if non_empty.count(header) > 1})
            if duplicates:
                raise RuntimeError(f"Encabezados duplicados en hoja {sheet_name}: {', '.join(duplicates)}")

            header_map = {header: idx for idx, header in enumerate(headers) if header}
            missing = [header for header, _ in EXPECTED_COLUMNS if normalize_header(header) not in header_map]
            if not missing:
                return sheet_name, header_map, row_idx
            if best_missing is None or len(missing) < len(best_missing):
                best_missing = missing
    missing_text = ", ".join(best_missing or [header for header, _ in EXPECTED_COLUMNS])
    raise RuntimeError("No se encontro una hoja con los 19 encabezados esperados. Faltantes: " + missing_text)


def read_rows(path: Path) -> tuple[str, list[tuple], int, list[dict]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name, header_map, header_row = find_source_sheet(workbook)
    ws = workbook[sheet_name]

    rows: list[tuple] = []
    fecha_errors: list[dict] = []
    empty_rows = 0

    for excel_row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        values = []
        raw_values = []
        for header, _sql_name in EXPECTED_COLUMNS:
            idx = header_map[normalize_header(header)]
            raw = row[idx] if idx < len(row) else None
            raw_values.append(raw)
            values.append(value_to_text(raw))

        if all(is_empty_value(raw) for raw in raw_values):
            empty_rows += 1
            continue

        fecha_raw = raw_values[[header for header, _ in EXPECTED_COLUMNS].index("FECH_GEST")]
        fecha_filtro = parse_fecha_gestion(fecha_raw)
        if fecha_filtro is None:
            fecha_errors.append(
                {
                    "numero_fila_origen": excel_row_idx,
                    "fech_gest": value_to_text(fecha_raw),
                    "mensaje": "FECH_GEST vacio o no convertible",
                }
            )
            continue

        rows.append((excel_row_idx, values, fecha_filtro))

    return sheet_name, rows, empty_rows, fecha_errors


def insert_rows(path: Path, sheet_name: str, rows: list[tuple], id_ejecucion: uuid.UUID, verbose: bool = True) -> int:
    insert_cols = [
        "id_ejecucion",
        "nombre_archivo",
        "numero_fila_origen",
        *[sql_name for _header, sql_name in EXPECTED_COLUMNS],
        "fecha_gestion_filtro",
    ]
    placeholders = ",".join("?" for _ in insert_cols)
    sql = f"INSERT INTO {TABLE} ({','.join(map(sql_ident, insert_cols))}) VALUES ({placeholders})"

    payload = []
    for numero_fila, values, fecha_filtro in rows:
        payload.append((str(id_ejecucion), path.name, numero_fila, *values, fecha_filtro))

    inserted = 0
    with connect() as cn:
        cn.autocommit = False
        cur = cn.cursor()
        cur.fast_executemany = True
        try:
            if verbose:
                print(f"Insertando en bloques de {BATCH_SIZE}...")
            for i in range(0, len(payload), BATCH_SIZE):
                batch = payload[i : i + BATCH_SIZE]
                cur.executemany(sql, batch)
                inserted += len(batch)
                block_number = (i // BATCH_SIZE) + 1
                if verbose:
                    print(f"Bloque {block_number}: {len(batch)} filas insertadas")
            cn.commit()
        except Exception:
            cn.rollback()
            raise
    return inserted


def success_result(id_ejecucion: uuid.UUID, path: Path, sheet_name: str, detected: int, empty: int, inserted: int) -> dict:
    return {
        "ok": True,
        "id_ejecucion": str(id_ejecucion),
        "archivo": path.name,
        "hoja": sheet_name,
        "filas_detectadas": detected,
        "filas_vacias_omitidas": empty,
        "filas_insertadas": inserted,
        "filas_fecha_no_convertible": 0,
        "filas_error": 0,
    }


def error_result(message: str, path: Path | None = None, details: list[dict] | None = None) -> dict:
    return {
        "ok": False,
        "archivo": path.name if path else None,
        "error": message,
        "filas_error": len(details or []),
        "detalle": details or [],
    }


def print_pretty_result(result: dict) -> None:
    print()
    print("Carga Gestiones Diarias SCT")
    print()
    print(f"Archivo: {result.get('archivo') or '-'}")
    if result.get("hoja"):
        print(f"Hoja: {result.get('hoja')}")
    if result.get("id_ejecucion"):
        print(f"ID ejecucion: {result.get('id_ejecucion')}")
    print()
    print(f"Filas detectadas: {result.get('filas_detectadas', 0)}")
    print(f"Filas vacias omitidas: {result.get('filas_vacias_omitidas', 0)}")
    print(f"Filas insertadas: {result.get('filas_insertadas', 0)}")
    print(f"Filas fecha no convertible: {result.get('filas_fecha_no_convertible', 0)}")
    print(f"Filas error: {result.get('filas_error', 0)}")
    if result.get("error"):
        print()
        print(f"Error: {result.get('error')}")
        details = result.get("detalle") or []
        if details:
            print("Detalle:")
            for item in details[:10]:
                print(f"- {item}")
            if len(details) > 10:
                print(f"- ... {len(details) - 10} errores adicionales")
    print()
    print(f"Resultado: {'OK' if result.get('ok') else 'ERROR'}")


def run(file_path: Path, verbose: bool = True) -> dict:
    validate_path(file_path)
    sheet_name, rows, empty_rows, fecha_errors = read_rows(file_path)
    if fecha_errors:
        return error_result("FECH_GEST vacio o no convertible", file_path, fecha_errors[:100])

    id_ejecucion = uuid.uuid4()
    inserted = insert_rows(file_path, sheet_name, rows, id_ejecucion, verbose)
    return success_result(id_ejecucion, file_path, sheet_name, len(rows) + empty_rows, empty_rows, inserted)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Carga gestiones diarias SCT en SQL Server")
    parser.add_argument("--file", help="Ruta del archivo Excel a procesar. Si se omite, usa la carpeta por defecto de Gestion SCT")
    parser.add_argument("--json", action="store_true", help="Imprime el resultado final como JSON para automatizaciones")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    path = resolve_input_file(args.file)
    logger = setup_logger()
    try:
        result = run(path, verbose=not args.json)
    except Exception as exc:
        result = error_result(str(exc), path)

    if result.get("ok"):
        logger.info(json.dumps(result, ensure_ascii=False))
    else:
        logger.error(json.dumps(result, ensure_ascii=False))

    if args.json:
        print(json.dumps(result, ensure_ascii=False))
    else:
        print_pretty_result(result)
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
