from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, SharedItems, WorksheetSource
from openpyxl.pivot.table import DataField, FieldItem, Location, PageField, PivotField, PivotTableStyle, RowColField, TableDefinition
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import run_query, run_query_sets
from services.la_araucana_service import (
    ASIGNACION_TABLE,
    CARTERA_CRM,
    EJECUTIVOS_TABLE,
    GESTION_TABLE,
    PAGOS_TABLE,
    RESPUESTA_RANK_TABLE,
    TIPOS_PAGO_VALIDOS,
    _atribucion_sql,
    _columns,
    _contacto_gestion_order_expr,
    _deuda_cte,
    _ejecutivas,
    _parse_period,
    _pick_optional,
    _resolved_cols,
    _to_mes_proceso,
)


CARTERAS = ["VIGENTE", "CASTIGO", "+365"]
FONT = "Arial"
MONEY = "#,##0"
PCT = "0.00%"
DATE = "dd-mm-yyyy"

HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="1F4E78")
TITLE_FONT = Font(name=FONT, size=14, bold=True)
BOLD = Font(name=FONT, size=10, bold=True)
NORMAL = Font(name=FONT, size=10)
TOTAL_FILL = PatternFill("solid", start_color="DDEBF7")

# Columnas de la hoja Pagos: (clave, encabezado, ancho, formato).
PAGOS_COLUMNS = [
    ("contrato", "Folio / contrato", 16, None),
    ("rut", "RUT", 12, None),
    ("tipo_cartera", "Tipo cartera", 13, None),
    ("fecha_pago", "Fecha pago", 12, DATE),
    ("tipo_pago", "Tipo pago", 14, None),
    ("recupero", "Recupero", 14, MONEY),
    ("ejecutivo", "Ejecutiva asignada", 22, None),
    ("criterio", "Criterio", 27, None),
    ("usuario", "Usuario gestion ganadora", 24, None),
    ("contacto", "Contacto", 20, None),
    ("respuesta", "Respuesta", 34, None),
    ("ranking", "Ranking respuesta", 17, None),
    ("fecha_gestion", "Fecha gestion", 14, DATE),
    ("hora_gestion", "Hora gestion", 13, None),
    ("telefono", "Telefono", 13, None),
    ("gestiones_antes", "Gestiones del RUT hasta el pago", 29, None),
    ("gestiones_despues", "Gestiones del RUT despues del pago", 32, None),
]


def _pagos_col(key: str) -> str:
    return get_column_letter([k for k, _h, _w, _f in PAGOS_COLUMNS].index(key) + 1)


def _cell(ws, value, font=NORMAL, fmt=None, fill=None, wrap=False):
    cell = WriteOnlyCell(ws, value=value)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def _data_sheet(wb, title, columns, rows, extra=()):
    """Hoja de detalle. `extra` agrega columnas con formula: (encabezado, ancho, funcion(fila_excel))."""
    ws = wb.create_sheet(title)
    ws.freeze_panes = "A2"
    widths = [width for _key, _header, width, _fmt in columns] + [width for _header, width, _fn in extra]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{len(rows) + 1}"
    headers = [header for _key, header, _width, _fmt in columns] + [header for header, _width, _fn in extra]
    ws.append([_cell(ws, header, HEADER_FONT, fill=HEADER_FILL) for header in headers])
    for number, row in enumerate(rows, start=2):
        cells = [_cell(ws, row.get(key), fmt=fmt) for key, _header, _width, fmt in columns]
        cells += [_cell(ws, fn(number)) for _header, _width, fn in extra]
        ws.append(cells)
    return ws


def _raw_value(ws, value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if value.startswith("="):
            # Texto que parte con "=": se guarda como texto, no como formula.
            cell = WriteOnlyCell(ws, value=value)
            cell.data_type = "s"
            return cell
    return value


def _sheet_gestiones(wb, columns: list[str], rows: list[dict]) -> None:
    """Gestiones tal como vienen en la tabla; solo cambia el orden de las filas."""
    ws = wb.create_sheet("Gestiones")
    ws.freeze_panes = "A2"
    first = rows[0] if rows else {}
    for idx, name in enumerate(columns, start=1):
        # Las columnas con fecha y hora necesitan mas ancho para no verse como ####.
        width = 20 if isinstance(first.get(name), datetime) else max(12, min(len(name) + 6, 30))
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    ws.append([_cell(ws, name, HEADER_FONT, fill=HEADER_FILL) for name in columns])
    for row in rows:
        ws.append([_raw_value(ws, row.get(name)) for name in columns])


def _table_columns(table_name: str) -> list[str]:
    schema, table = table_name.split(".", 1)
    rows = run_query(
        """
        SELECT c.name
        FROM sys.columns c
        INNER JOIN sys.tables t ON c.object_id = t.object_id
        INNER JOIN sys.schemas s ON t.schema_id = s.schema_id
        WHERE s.name = ? AND t.name = ?
        ORDER BY c.column_id
        """,
        (schema, table),
    )
    return [r["name"] for r in rows]


def _fetch(periodo: str) -> dict:
    c = _resolved_cols()
    if not c["id_gest"]:
        raise RuntimeError("La tabla de gestiones no tiene columna id; no se puede generar la auditoría.")
    ejecutivas = _ejecutivas(periodo)
    mes_proceso = _to_mes_proceso(periodo)
    _period_month, _period_day, month_start, month_end, _file_tokens = _parse_period(mes_proceso)
    base_sql, base_params, period_month = _atribucion_sql(c, periodo)
    deuda_sql, deuda_params = _deuda_cte(periodo)
    # Orden de la hoja Gestiones: dentro de cada RUT, primero la gestion que gana segun la regla,
    # asi un BUSCARV por RUT devuelve la gestion ganadora.
    posterior = "(p.rut IS NOT NULL AND t.fecha_gestion > p.fecha_pago)"
    sql = f"""{base_sql}
    SELECT
        b.contrato,
        b.rut,
        b.tipo_cartera,
        b.fecha_pago,
        b.tipo_pago,
        b.recupero,
        COALESCE(b.ejecutivo, 'SIN GESTION') AS ejecutivo,
        b.criterio,
        b.usuario,
        b.contacto,
        b.respuesta,
        CASE WHEN b.respuesta_ranking = 999999 THEN NULL ELSE b.respuesta_ranking END AS ranking,
        b.fecha_gestion,
        b.hora_gestion,
        b.telefono,
        COALESCE(n.gestiones_antes, 0) AS gestiones_antes,
        COALESCE(n.gestiones_despues, 0) AS gestiones_despues
    FROM #base b
    LEFT JOIN (
        SELECT
            p.pago_id,
            SUM(CASE WHEN g.fecha_gestion <= p.fecha_pago THEN 1 ELSE 0 END) AS gestiones_antes,
            SUM(CASE WHEN g.fecha_gestion > p.fecha_pago THEN 1 ELSE 0 END) AS gestiones_despues
        FROM #pagos p
        INNER JOIN #gest g ON g.rut = p.rut
        GROUP BY p.pago_id
    ) n ON n.pago_id = b.pago_id
    ORDER BY b.tipo_cartera, b.ejecutivo, b.fecha_pago, b.rut;

    SELECT g.*
    FROM {GESTION_TABLE} g
    LEFT JOIN #gest t ON t.id_gestion = g.{c['id_gest']}
    LEFT JOIN (
        SELECT rut, MIN(fecha_pago) AS fecha_pago
        FROM #pagos
        WHERE fecha_pago IS NOT NULL
        GROUP BY rut
    ) p ON p.rut = CONVERT(varchar(50), g.{c['rut_gest']})
    WHERE g.{c['cartera_gest']} = {CARTERA_CRM}
      AND g.{c['fecha_gest']} >= CAST(? AS date)
      AND g.{c['fecha_gest']} <= CAST(? AS date)
    ORDER BY
        g.{c['rut_gest']},
        CASE WHEN t.id_gestion IS NULL THEN 1 ELSE 0 END,
        CASE WHEN {posterior} THEN 1 ELSE 0 END,
        CASE WHEN {posterior} THEN 0 ELSE t.respuesta_ranking END,
        CASE WHEN {posterior} THEN 0 ELSE {_contacto_gestion_order_expr("t.contacto")} END,
        CASE WHEN {posterior} THEN NULL ELSE t.fecha_gestion END DESC,
        CASE WHEN {posterior} THEN NULL ELSE t.hora_gestion END DESC,
        CASE WHEN {posterior} THEN NULL ELSE t.id_gestion END DESC,
        t.fecha_gestion,
        t.hora_gestion,
        t.id_gestion;

    WITH {deuda_sql.lstrip()}
    SELECT d.folio, d.rut, d.tipo_cartera, d.deuda, COALESCE(d.ejecutivo, 'SIN GESTION') AS ejecutivo, d.criterio, d.usuario_mejor_gestion
    FROM deuda_folio d
    ORDER BY d.tipo_cartera, d.ejecutivo, d.folio;
    """
    params = tuple(base_params + [month_start, month_end] + deuda_params)
    pagos, gestiones, deuda = run_query_sets(sql, params)[-3:]

    rank_cols = _columns(RESPUESTA_RANK_TABLE)
    rank_resp = _pick_optional(rank_cols, ["Respuesta", "respuesta", "RESPUESTA", "RespuestaGestion", "respuesta_gestion"])
    rank_num = _pick_optional(rank_cols, ["RANKING", "ranking", "Ranking"])
    ranking = (
        run_query(f"SELECT {rank_num} AS ranking, {rank_resp} AS respuesta FROM {RESPUESTA_RANK_TABLE} ORDER BY {rank_num}, {rank_resp}")
        if rank_resp and rank_num
        else []
    )
    return {
        "period_month": period_month,
        "mes_proceso": mes_proceso,
        "ejecutivas": ejecutivas,
        "pagos": pagos,
        "gestiones": gestiones,
        "gestiones_columns": _table_columns(GESTION_TABLE),
        "rut_gest": c["rut_gest"],
        "usuario_gest": c["usuario_gest"],
        "deuda": deuda,
        "ranking": ranking,
    }


def _sheet_explicacion(wb, data: dict) -> None:
    ws = wb.create_sheet("Como se calcula")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    n_pagos = len(data["pagos"]) + 1
    n_deuda = len(data["deuda"]) + 1
    rec = _pagos_col("recupero")
    crit = _pagos_col("criterio")

    def line(a="", b="", font_a=BOLD, fmt=None):
        ws.append([_cell(ws, a, font_a), _cell(ws, b, fmt=fmt, wrap=True)])

    ws.append([_cell(ws, "Auditoria Caja La Araucana: como se asocia cada pago a una ejecutiva", TITLE_FONT)])
    line("Mes de proceso", data["mes_proceso"])
    line("Generado", datetime.now().strftime("%d-%m-%Y %H:%M"))
    line()
    line("Fuentes de datos", "")
    line("Pagos", f"{PAGOS_TABLE}: recuperacion, fecha de pago, tipo de pago, contrato y RUT del afiliado.")
    line("Gestiones", f"{GESTION_TABLE}, cartera CRM {CARTERA_CRM}: usuario, contacto, respuesta, fecha y hora.")
    line("Ranking de respuestas", f"{RESPUESTA_RANK_TABLE}: 1 es la mejor respuesta. Ver hoja 'Ranking'.")
    line("Ejecutivas", f"{EJECUTIVOS_TABLE}, cartera {CARTERA_CRM}, vigentes en el mes. Ver hoja 'Ejecutivas'.")
    line("Deuda y tipo de cartera", f"{ASIGNACION_TABLE} del mismo mes de proceso.")
    line()
    line("Reglas", "")
    line("1. Pagos que cuentan", "Pagos del mes de proceso con tipo de pago: " + ", ".join(TIPOS_PAGO_VALIDOS) + ". Cada pago es una fila de la hoja 'Pagos'.")
    line("2. Gestiones que cuentan", "Solo las gestiones con fecha dentro del mes seleccionado y con usuario informado. Estan todas en la hoja 'Gestiones'.")
    line("3. Cruce", "El pago se cruza con las gestiones por RUT (el RUT del afiliado sin digito verificador). La gestion no trae folio.")
    line(
        "4. Gestion que gana el pago",
        "Entre las gestiones del RUT con fecha menor o igual a la fecha de pago gana la de mejor ranking de respuesta "
        "(las respuestas sin ranking quedan al final). A igual ranking, el mejor contacto: Contacto directo, Contacto indirecto, "
        "No contactado, Gestion discador. A igualdad, la mas cercana al pago (fecha y hora mas recientes). "
        "Una gestion del mismo dia del pago cuenta como anterior, porque el pago no trae hora.",
    )
    line("5. Sin gestion antes del pago", "Si el RUT no tiene gestiones hasta la fecha de pago, gana la primera gestion posterior al pago dentro del mes.")
    line("6. Sin ninguna gestion en el mes", "El pago queda como 'SIN GESTION' y no entra al resumen.")
    line("7. Nombre de la ejecutiva", f"El usuario de la gestion ganadora se busca en {EJECUTIVOS_TABLE}. Si no esta, el pago se agrupa en PHOENIX.")
    line("8. Tipo de cartera", "Es el de la asignacion del folio (Vigente, Castigo, +365). Si el folio no esta en la asignacion, el del archivo de pagos.")
    line(
        "9. Deuda",
        "Cada folio asignado va a la ejecutiva que se llevo su pago mas reciente. Si el folio no tiene pago, va a la mejor gestion "
        "del RUT en el mes (ranking, contacto y la mas reciente). Sin gestion en el mes, la deuda no entra al resumen. Ver hoja 'Deuda'.",
    )
    line("10. Cumplimiento por cartera", "Recupero de la ejecutiva en la cartera dividido por el recupero total de esa cartera.")
    line("11. Aporte final", "Recupero de la ejecutiva en las tres carteras dividido por el recupero total de las tres carteras.")
    line()
    line("Como leer las hojas", "")
    line("Resumen", "Lo mismo que muestra la pagina, calculado con formulas sobre las hojas 'Pagos' y 'Deuda'.")
    line("Tabla dinamica", "Tabla dinamica sobre la hoja 'Pagos': recupero por ejecutiva y tipo de cartera, con filtro por criterio. Se puede reordenar.")
    line("Pagos", "Un pago por fila, con la ejecutiva asignada, el criterio usado y la gestion que lo gano.")
    line(
        "Gestiones",
        f"Todas las gestiones del mes tal como vienen en {GESTION_TABLE}, sin columnas agregadas. Solo cambia el orden: dentro de cada RUT "
        "la primera fila es la gestion que gana segun la regla (para un RUT con pago, respecto de su fecha de pago; para un RUT sin pago, "
        "la mejor gestion del mes). Por eso un BUSCARV por RUT contra esta hoja devuelve la gestion ganadora.",
    )
    line(
        "BUSCARV",
        "En la hoja 'Pagos', la columna 'Usuario segun BUSCARV' busca el RUT del pago en la hoja 'Gestiones' y la columna 'Coincide' lo compara "
        "con el usuario asignado. Solo puede dar NO cuando un mismo RUT tiene pagos en fechas distintas que ganan gestiones distintas: "
        "la hoja queda ordenada segun la primera fecha de pago del RUT.",
    )
    line("Deuda", "Un folio asignado por fila, con la ejecutiva a la que se le cuenta la deuda y por que.")
    line()
    line("Cuadre", "")
    ws.append([_cell(ws, "Pagos validos"), _cell(ws, f"=COUNTA(Pagos!$A$2:$A${n_pagos})", fmt=MONEY)])
    ws.append([_cell(ws, "Recupero de pagos validos"), _cell(ws, f"=SUM(Pagos!${rec}$2:${rec}${n_pagos})", fmt=MONEY)])
    for label in ("GESTION ANTES DEL PAGO", "GESTION DESPUES DEL PAGO", "SIN GESTION"):
        ws.append([
            _cell(ws, f"Recupero: {label.lower()}"),
            _cell(ws, f'=SUMIF(Pagos!${crit}$2:${crit}${n_pagos},"{label}",Pagos!${rec}$2:${rec}${n_pagos})', fmt=MONEY),
        ])
    ws.append([
        _cell(ws, "Recupero en el resumen", BOLD),
        _cell(ws, f'=SUM(Pagos!${rec}$2:${rec}${n_pagos})-SUMIF(Pagos!${crit}$2:${crit}${n_pagos},"SIN GESTION",Pagos!${rec}$2:${rec}${n_pagos})', BOLD, MONEY),
    ])
    ws.append([_cell(ws, "Deuda asignada total"), _cell(ws, f"=SUM(Deuda!$D$2:$D${n_deuda})", fmt=MONEY)])
    ws.append([
        _cell(ws, "Deuda en el resumen", BOLD),
        _cell(ws, f'=SUM(Deuda!$D$2:$D${n_deuda})-SUMIF(Deuda!$E$2:$E${n_deuda},"SIN GESTION",Deuda!$D$2:$D${n_deuda})', BOLD, MONEY),
    ])
    ws.append([
        _cell(ws, "Pagos donde BUSCARV no coincide"),
        _cell(ws, f'=COUNTIF(Pagos!${get_column_letter(len(PAGOS_COLUMNS) + 2)}$2:${get_column_letter(len(PAGOS_COLUMNS) + 2)}${n_pagos},"NO")', fmt=MONEY),
    ])


def _sheet_resumen(wb, data: dict) -> None:
    ws = wb.create_sheet("Resumen")
    n_pagos = len(data["pagos"]) + 1
    n_deuda = len(data["deuda"]) + 1
    p_cartera, p_ejecutivo, p_recupero = _pagos_col("tipo_cartera"), _pagos_col("ejecutivo"), _pagos_col("recupero")
    nombres = sorted({r["ejecutivo"] for r in data["pagos"] + data["deuda"]} - {"SIN GESTION", "PHOENIX"})
    nombres.append("PHOENIX")
    ws.column_dimensions["A"].width = 26
    for idx in range(2, 2 + len(CARTERAS) * 3 + 2):
        ws.column_dimensions[get_column_letter(idx)].width = 17
    ws.freeze_panes = "B4"

    ws.append([_cell(ws, f"Resumen por ejecutiva - {data['mes_proceso']}", TITLE_FONT)])
    # Fila 2: nombre de la cartera sobre cada grupo; las formulas comparan contra estas celdas.
    fila_carteras = [_cell(ws, "", HEADER_FONT, fill=HEADER_FILL)]
    fila_titulos = [_cell(ws, "Ejecutivo", HEADER_FONT, fill=HEADER_FILL)]
    for cartera in CARTERAS:
        fila_carteras += [_cell(ws, cartera, HEADER_FONT, fill=HEADER_FILL), _cell(ws, "", HEADER_FONT, fill=HEADER_FILL), _cell(ws, "", HEADER_FONT, fill=HEADER_FILL)]
        fila_titulos += [_cell(ws, label, HEADER_FONT, fill=HEADER_FILL) for label in ("Deuda", "Recupero", "Cumplimiento")]
    fila_carteras += [_cell(ws, "", HEADER_FONT, fill=HEADER_FILL), _cell(ws, "", HEADER_FONT, fill=HEADER_FILL)]
    fila_titulos += [_cell(ws, "Recupero total", HEADER_FONT, fill=HEADER_FILL), _cell(ws, "Aporte final", HEADER_FONT, fill=HEADER_FILL)]
    ws.append(fila_carteras)
    ws.append(fila_titulos)

    first = 4
    total_row = first + len(nombres)
    col_total = get_column_letter(2 + len(CARTERAS) * 3)
    for offset, nombre in enumerate(nombres):
        r = first + offset
        row = [_cell(ws, nombre)]
        recuperos = []
        for idx, _cartera in enumerate(CARTERAS):
            col_deuda = get_column_letter(2 + idx * 3)
            col_rec = get_column_letter(3 + idx * 3)
            recuperos.append(f"{col_rec}{r}")
            row.append(_cell(ws, f"=SUMPRODUCT((Deuda!$C$2:$C${n_deuda}={col_deuda}$2)*(Deuda!$E$2:$E${n_deuda}=$A{r})*Deuda!$D$2:$D${n_deuda})", fmt=MONEY))
            row.append(_cell(
                ws,
                f"=SUMPRODUCT((Pagos!${p_cartera}$2:${p_cartera}${n_pagos}={col_deuda}$2)*(Pagos!${p_ejecutivo}$2:${p_ejecutivo}${n_pagos}=$A{r})*Pagos!${p_recupero}$2:${p_recupero}${n_pagos})",
                fmt=MONEY,
            ))
            row.append(_cell(ws, f"=IF({col_rec}${total_row}=0,0,{col_rec}{r}/{col_rec}${total_row})", fmt=PCT))
        row.append(_cell(ws, "=" + "+".join(recuperos), fmt=MONEY))
        row.append(_cell(ws, f"=IF(${col_total}${total_row}=0,0,{col_total}{r}/${col_total}${total_row})", fmt=PCT))
        ws.append(row)

    total = [_cell(ws, "Total general", BOLD, fill=TOTAL_FILL)]
    for idx in range(len(CARTERAS) * 3 + 2):
        col = get_column_letter(2 + idx)
        es_pct = (idx < len(CARTERAS) * 3 and idx % 3 == 2) or idx == len(CARTERAS) * 3 + 1
        total.append(_cell(ws, f"=SUM({col}{first}:{col}{total_row - 1})", BOLD, PCT if es_pct else MONEY, TOTAL_FILL))
    ws.append(total)
    ws.append([])
    ws.append([_cell(ws, "Los pagos y folios marcados 'SIN GESTION' no entran a este resumen. El cuadre esta en la hoja 'Como se calcula'.")])


def _sheet_tabla_dinamica(wb, data: dict, headers: list[str]) -> None:
    """Tabla dinamica sobre la hoja Pagos; Excel la calcula al abrir el archivo."""
    ws = wb.create_sheet("Tabla dinamica")
    ws.column_dimensions["A"].width = 30
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18
    ws.append([_cell(ws, f"Recupero por ejecutiva y tipo de cartera - {data['mes_proceso']}", TITLE_FONT)])

    keys = [k for k, _h, _w, _f in PAGOS_COLUMNS]
    i_fila, i_col, i_filtro, i_valor = keys.index("ejecutivo"), keys.index("tipo_cartera"), keys.index("criterio"), keys.index("recupero")
    last_col = get_column_letter(len(headers))
    cache = CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(ref=f"A1:{last_col}{len(data['pagos']) + 1}", sheet="Pagos")),
        cacheFields=[CacheField(name=header, sharedItems=SharedItems()) for header in headers],
        refreshOnLoad=True,
        saveData=False,
        recordCount=0,
        createdVersion=5,
        refreshedVersion=5,
        minRefreshableVersion=3,
    )
    fields = []
    for idx in range(len(headers)):
        if idx == i_fila:
            fields.append(PivotField(axis="axisRow", showAll=False, items=[FieldItem(t="default")]))
        elif idx == i_col:
            fields.append(PivotField(axis="axisCol", showAll=False, items=[FieldItem(t="default")]))
        elif idx == i_filtro:
            fields.append(PivotField(axis="axisPage", showAll=False, items=[FieldItem(t="default")]))
        elif idx == i_valor:
            fields.append(PivotField(dataField=True, showAll=False))
        else:
            fields.append(PivotField(showAll=False))
    pivot = TableDefinition(
        name="RecuperoPorEjecutiva",
        cacheId=1,
        dataCaption="Valores",
        createdVersion=5,
        updatedVersion=5,
        minRefreshableVersion=3,
        useAutoFormatting=True,
        itemPrintTitles=True,
        indent=0,
        outline=True,
        outlineData=True,
        multipleFieldFilters=False,
        applyNumberFormats=False,
        applyBorderFormats=False,
        applyFontFormats=False,
        applyPatternFormats=False,
        applyAlignmentFormats=False,
        applyWidthHeightFormats=True,
        location=Location(ref="A5:E12", firstHeaderRow=1, firstDataRow=2, firstDataCol=1, rowPageCount=1, colPageCount=1),
        pivotFields=fields,
        rowFields=[RowColField(x=i_fila)],
        colFields=[RowColField(x=i_col)],
        pageFields=[PageField(fld=i_filtro, hier=-1)],
        dataFields=[DataField(name="Suma de Recupero", fld=i_valor, baseField=0, baseItem=0, numFmtId=3)],
        pivotTableStyleInfo=PivotTableStyle(name="PivotStyleMedium2", showRowHeaders=True, showColHeaders=True, showRowStripes=False, showColStripes=False, showLastColumn=True),
    )
    pivot.cache = cache
    ws._pivots.append(pivot)


def build_auditoria_workbook(periodo: str) -> tuple[str, Workbook]:
    data = _fetch(periodo)
    wb = Workbook(write_only=True)
    _sheet_explicacion(wb, data)
    _sheet_resumen(wb, data)

    # BUSCARV del RUT del pago contra la hoja Gestiones (desde la columna del RUT hasta la del usuario).
    gest_cols = data["gestiones_columns"]
    i_rut, i_usuario = gest_cols.index(data["rut_gest"]), gest_cols.index(data["usuario_gest"])
    extra = []
    if i_usuario > i_rut:
        rango = f"Gestiones!${get_column_letter(i_rut + 1)}:${get_column_letter(i_usuario + 1)}"
        rut, usuario = _pagos_col("rut"), _pagos_col("usuario")
        buscarv = get_column_letter(len(PAGOS_COLUMNS) + 1)
        extra = [
            ("Usuario segun BUSCARV", 24, lambda n: f'=IFERROR(VLOOKUP({rut}{n},{rango},{i_usuario - i_rut + 1},FALSE)&"","")'),
            ("Coincide", 10, lambda n: f'=IF(UPPER(TRIM({buscarv}{n}))=UPPER(TRIM({usuario}{n}&"")),"SI","NO")'),
        ]
    headers = [header for _key, header, _width, _fmt in PAGOS_COLUMNS] + [header for header, _width, _fn in extra]
    _sheet_tabla_dinamica(wb, data, headers)
    _data_sheet(wb, "Pagos", PAGOS_COLUMNS, data["pagos"], extra)
    _sheet_gestiones(wb, gest_cols, data["gestiones"])
    _data_sheet(
        wb,
        "Deuda",
        [
            ("folio", "Folio", 16, None),
            ("rut", "RUT", 12, None),
            ("tipo_cartera", "Tipo cartera", 13, None),
            ("deuda", "Deuda total", 15, MONEY),
            ("ejecutivo", "Ejecutiva asignada", 22, None),
            ("criterio", "Criterio", 24, None),
            ("usuario_mejor_gestion", "Usuario mejor gestion del mes", 28, None),
        ],
        data["deuda"],
    )
    _data_sheet(wb, "Ranking", [("ranking", "Ranking", 10, None), ("respuesta", "Respuesta", 42, None)], data["ranking"])
    _data_sheet(
        wb,
        "Ejecutivas",
        [("usuario", "Usuario CRM", 18, None), ("nombre", "Nombre", 28, None)],
        [{"usuario": usuario, "nombre": nombre} for usuario, nombre in sorted(data["ejecutivas"].items())],
    )
    return data["period_month"], wb
