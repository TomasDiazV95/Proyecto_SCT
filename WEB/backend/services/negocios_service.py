"""Negocios cerrados del mes: carga de la plantilla de cada negocio y conteo por ejecutivo.

Las administrativas suben un Excel por negocio (Panel Administrativo > Negocios) con las columnas
OP, TIPO NEGOCIO, EJECUTIVO y PERIODO (mm-yyyy); cada productividad muestra el conteo del mes
en su pestaña Negocios.
"""
from __future__ import annotations

import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import cache
from database import get_connection, run_query
from services.sc_temprana_service import DERIVADO_TO_USER, USER_TO_NAME


CARGAS = "dbo.negocios_cargas"
CIERRE = "dbo.negocios_cierre"
MAX_BYTES = 10 * 1024 * 1024
# El encabezado no siempre esta en la primera fila de la hoja.
FILAS_ENCABEZADO = 10
SIN_EJECUTIVO = "Sin ejecutivo"
SIN_TIPO = "SIN TIPO"
# Valor de la columna tipo de las tablas: cada negocio sube un solo archivo.
TIPO_CARGA = "negocios"

NEGOCIOS = {"sc-terreno": "SC Terreno", "sc-telefonia": "SC Telefonia", "gm": "GM"}

# Columnas de la plantilla, en orden. Los encabezados se comparan normalizados (mayusculas, sin
# acentos ni simbolos) y se aceptan las variantes indicadas.
PLANTILLA = ["OP", "TIPO NEGOCIO", "EJECUTIVO", "PERIODO"]
COLUMNAS = {
    "operacion": ("OP", "OPERACION"),
    "producto": ("TIPO NEGOCIO", "TIPO DE NEGOCIO"),
    "ejecutivo": ("EJECUTIVO",),
    "periodo": ("PERIODO",),
}

PRODUCTO_ORDEN = ["RECONDUCCION", "REFINANCIAMIENTO", "NOVACION", "DACION", "EXTENSION"]
PRODUCTO_NOMBRE = {
    "RECONDUCCION": "Reconducciones",
    "REFINANCIAMIENTO": "Refinanciamientos",
    "NOVACION": "Novaciones",
    "DACION": "Daciones",
    "EXTENSION": "Extensiones",
}
# Como empieza cada tipo de negocio: 'Reco', 'RECONDUCCIONES' y 'Reconducción' son el mismo.
PRODUCTO_PREFIJO = {
    "RECO": "RECONDUCCION",
    "REFI": "REFINANCIAMIENTO",
    "NOVA": "NOVACION",
    "DACI": "DACION",
    "EXTEN": "EXTENSION",
}

# negocios_cargas: una fila por negocio y mes cargado. negocios_cierre: una por operacion.
SQL_TABLAS = f"""
IF OBJECT_ID('{CARGAS}', 'U') IS NULL
BEGIN
    CREATE TABLE {CARGAS} (
        id_carga INT IDENTITY(1,1) NOT NULL CONSTRAINT PK_negocios_cargas PRIMARY KEY,
        negocio VARCHAR(20) NOT NULL,
        tipo VARCHAR(20) NOT NULL,
        periodo DATE NOT NULL,
        nombre_archivo NVARCHAR(260) NOT NULL,
        filas INT NOT NULL,
        cargado_por NVARCHAR(200) NULL,
        fecha_carga DATETIME2 NOT NULL CONSTRAINT DF_negocios_cargas_fecha DEFAULT (SYSDATETIME()),
        CONSTRAINT UQ_negocios_cargas UNIQUE (negocio, tipo, periodo)
    );
END;

IF OBJECT_ID('{CIERRE}', 'U') IS NULL
BEGIN
    CREATE TABLE {CIERRE} (
        id BIGINT IDENTITY(1,1) NOT NULL CONSTRAINT PK_negocios_cierre PRIMARY KEY,
        id_carga INT NOT NULL CONSTRAINT FK_negocios_cierre_carga REFERENCES {CARGAS}(id_carga),
        negocio VARCHAR(20) NOT NULL,
        tipo VARCHAR(20) NOT NULL,
        periodo DATE NOT NULL,
        producto VARCHAR(40) NOT NULL,
        operacion VARCHAR(40) NULL,
        rut VARCHAR(20) NULL,
        estado NVARCHAR(100) NULL,
        ejecutivo NVARCHAR(120) NOT NULL,
        ejecutivo_origen NVARCHAR(120) NULL,
        ciclo VARCHAR(20) NULL,
        empresa NVARCHAR(100) NULL,
        cliente NVARCHAR(200) NULL
    );

    CREATE INDEX IX_negocios_cierre_negocio_periodo ON {CIERRE} (negocio, periodo);
END;
"""

_tablas_listas = False


class ArchivoInvalido(ValueError):
    pass


def _ensure_tables() -> None:
    global _tablas_listas
    if _tablas_listas:
        return
    with get_connection() as cn:
        cn.cursor().execute(SQL_TABLAS)
        cn.commit()
    _tablas_listas = True


def plantilla(negocio: str) -> tuple[str, BytesIO]:
    """(nombre de archivo, Excel vacio con las columnas de la plantilla) del negocio."""
    if negocio not in NEGOCIOS:
        raise ArchivoInvalido(f"Negocio inválido: {negocio}")
    wb = Workbook()
    ws = wb.active
    ws.title = "NEGOCIOS"
    ws.append(PLANTILLA)
    for idx, cell in enumerate(ws[1], start=1):
        cell.font = Font(bold=True)
        columna = ws.column_dimensions[get_column_letter(idx)]
        columna.width = 22
        # Como texto: Excel no convierte '09-2026' en fecha ni la OP en notacion cientifica.
        columna.number_format = "@"
    ws.freeze_panes = "A2"
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return f"Plantilla_Negocios_{NEGOCIOS[negocio].replace(' ', '_')}.xlsx", output


def _norm(value) -> str:
    """Mayusculas, sin acentos y con espacios simples: para comparar nombres."""
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return " ".join(text.upper().split())


def _encabezado(value) -> str:
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", _norm(value)).split())


def _texto(value) -> str:
    """Valor de celda como texto; las operaciones vienen como numero."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _mes(periodo: str | None) -> str:
    """Mes 'YYYY-MM' desde cualquier formato de periodo de las productividades
    ('YYYY-MM', 'YYYY-MM-DD', 'YYYYMMDD', con o sin '(Cierre)')."""
    match = re.search(r"(20\d{2})\D?(0[1-9]|1[0-2])", str(periodo or ""))
    if not match:
        raise ArchivoInvalido(f"Periodo inválido: {periodo}. Se espera YYYY-MM")
    return f"{match.group(1)}-{match.group(2)}"


def _mes_plantilla(value) -> str | None:
    """Mes 'YYYY-MM' de la columna PERIODO: 'mm-yyyy' (tambien 'mm/yyyy' o 'yyyy-mm'), o la fecha
    en que Excel convierte ese texto. None si no se entiende."""
    if isinstance(value, (datetime, date)):
        return f"{value.year}-{value.month:02d}"
    text = _texto(value)
    match = re.fullmatch(r"(\d{1,2})\s*[-/.]\s*(20\d{2})", text)
    if match:
        mes, anio = int(match.group(1)), match.group(2)
    else:
        match = re.fullmatch(r"(20\d{2})\s*[-/.]\s*(\d{1,2})", text)
        if not match:
            return None
        anio, mes = match.group(1), int(match.group(2))
    return f"{anio}-{mes:02d}" if 1 <= mes <= 12 else None


def _leer_hoja(wb) -> list[dict] | None:
    """Filas de la primera hoja que tenga las columnas de la plantilla, con su numero de fila;
    None si ninguna las tiene."""
    for ws in wb.worksheets:
        filas = ws.iter_rows(values_only=True)
        indices: dict[str, int] | None = None
        numero = 0
        for _ in range(FILAS_ENCABEZADO):
            fila = next(filas, None)
            if fila is None:
                break
            numero += 1
            encabezados = {_encabezado(valor): idx for idx, valor in enumerate(fila) if valor is not None}
            encontrados = {
                campo: next((encabezados[v] for v in variantes if v in encabezados), None)
                for campo, variantes in COLUMNAS.items()
            }
            if all(idx is not None for idx in encontrados.values()):
                indices = encontrados
                break
        if indices is None:
            continue

        datos: list[dict] = []
        for fila in filas:
            numero += 1
            item = {campo: fila[idx] if idx < len(fila) else None for campo, idx in indices.items()}
            if any(_texto(valor) for valor in item.values()):
                datos.append({**item, "fila": numero})
        return datos
    return None


def _producto(value) -> str:
    clave = _norm(value)
    if not clave:
        return SIN_TIPO
    return next((p for prefijo, p in PRODUCTO_PREFIJO.items() if clave.startswith(prefijo)), clave)


def _ejecutivos_conocidos(negocio: str, mes: str) -> list[str]:
    """Ejecutivos del mes tal como los muestra la productividad del negocio."""
    if negocio == "sc-telefonia":
        return list(USER_TO_NAME.values())
    if negocio == "sc-terreno":
        # La sabana de SC Tardia trae al ejecutivo en mayusculas.
        rows = run_query(
            """
            SELECT DISTINCT UPPER(LTRIM(RTRIM(ejecutivo))) AS ejecutivo
            FROM dbo.stc_bloques_ejecutivos
            WHERE periodo = CAST(? AS date) AND activo = 1
            """,
            (f"{mes}-01",),
        )
    else:
        # Si el mes aun no tiene carterizado, todos los conocidos.
        sql = """
            SELECT DISTINCT LTRIM(RTRIM(ejecutivo)) AS ejecutivo
            FROM dbo.tmp_carterizado_GM
            WHERE ejecutivo IS NOT NULL AND LTRIM(RTRIM(ejecutivo)) <> ''
        """
        rows = run_query(sql + " AND mes_carterizado = CAST(? AS date)", (f"{mes}-01",)) or run_query(sql)
    return [r["ejecutivo"] for r in rows if r.get("ejecutivo")]


def _cruzar(nombre: str, conocidos: list[str]) -> str | None:
    """Ejecutivo conocido para el nombre del Excel: igual sin acentos, o el unico cuyas palabras
    estan contenidas ('Maria Yesenia Salinas' es 'Yesenia Salinas'; 'Lissette' es 'Lissette Rojas')."""
    clave = _norm(nombre)
    for conocido in conocidos:
        if _norm(conocido) == clave:
            return conocido
    palabras = set(clave.split())
    parecidos = [c for c in conocidos if set(_norm(c).split()) <= palabras or palabras <= set(_norm(c).split())]
    return parecidos[0] if len(parecidos) == 1 else None


def _resolver_ejecutivo(negocio: str, origen: str, conocidos: list[str]) -> tuple[str, bool]:
    """(nombre como lo muestra la productividad, si cruzo con un ejecutivo del mes)."""
    if not origen:
        return SIN_EJECUTIVO, True
    if negocio == "sc-telefonia":
        # Tambien se acepta el apodo ('ELY') o el usuario de gestion ('EMUNOZ').
        usuario = DERIVADO_TO_USER.get(_norm(origen)) or _norm(origen)
        if usuario in USER_TO_NAME:
            return USER_TO_NAME[usuario], True
    cruce = _cruzar(origen, conocidos)
    if cruce:
        return cruce, True
    return (origen.upper() if negocio == "sc-terreno" else origen), False


def _filas_con_error(filas: list[dict]) -> str:
    numeros = [str(f["fila"]) for f in filas]
    return ", ".join(numeros[:15]) + (" y otras" if len(numeros) > 15 else "")


def _leer_archivo(negocio: str, contenido: bytes) -> tuple[list[dict], dict]:
    try:
        wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:
        raise ArchivoInvalido("No se pudo leer el archivo. Debe ser un Excel .xlsx") from exc
    try:
        crudas = _leer_hoja(wb)
    finally:
        wb.close()
    if crudas is None:
        raise ArchivoInvalido(
            f"El archivo no corresponde a la plantilla: ninguna hoja tiene las columnas {', '.join(PLANTILLA)}"
        )
    if not crudas:
        raise ArchivoInvalido("El archivo no trae negocios")

    sin_periodo = [f for f in crudas if not _mes_plantilla(f["periodo"])]
    if sin_periodo:
        raise ArchivoInvalido(
            f"PERIODO vacío o inválido (se espera mm-yyyy, por ejemplo 09-2026) en las filas: {_filas_con_error(sin_periodo)}"
        )
    sin_op = [f for f in crudas if not _texto(f["operacion"])]
    if sin_op:
        raise ArchivoInvalido(f"Falta la OP en las filas: {_filas_con_error(sin_op)}")

    conocidos: dict[str, list[str]] = {}
    sin_cruce: set[str] = set()
    filas: list[dict] = []
    for cruda in crudas:
        mes = _mes_plantilla(cruda["periodo"])
        if mes not in conocidos:
            conocidos[mes] = _ejecutivos_conocidos(negocio, mes)
        origen = _texto(cruda["ejecutivo"])
        ejecutivo, cruza = _resolver_ejecutivo(negocio, origen, conocidos[mes])
        # Sin ejecutivos del mes en la productividad no hay contra que cruzar.
        if not cruza and conocidos[mes]:
            sin_cruce.add(origen)
        filas.append(
            {
                "periodo": mes,
                "operacion": _texto(cruda["operacion"]),
                "producto": _producto(cruda["producto"]),
                "ejecutivo": ejecutivo,
                "ejecutivo_origen": origen,
            }
        )
    avisos = {
        "sin_cruce": sorted(sin_cruce),
        "sin_tipo": sum(1 for f in filas if f["producto"] == SIN_TIPO),
        "sin_ejecutivo": sum(1 for f in filas if f["ejecutivo"] == SIN_EJECUTIVO),
    }
    return filas, avisos


def _carga_dict(row: dict) -> dict:
    return {
        "id_carga": row["id_carga"],
        "negocio": row["negocio"],
        "periodo": str(row["periodo"])[:7],
        "nombre_archivo": row["nombre_archivo"],
        "filas": int(row["filas"] or 0),
        "cargado_por": row.get("cargado_por") or "",
        "fecha_carga": row["fecha_carga"].isoformat(timespec="minutes") if row.get("fecha_carga") else "",
    }


def listar_cargas() -> list[dict]:
    _ensure_tables()
    rows = run_query(
        f"""
        SELECT id_carga, negocio, periodo, nombre_archivo, filas, cargado_por, fecha_carga
        FROM {CARGAS}
        ORDER BY periodo DESC, negocio
        """
    )
    return [_carga_dict(r) for r in rows]


def cargar(negocio: str, nombre_archivo: str, contenido: bytes, usuario: str) -> dict:
    """Guarda los negocios del archivo. Cada mes que trae (columna PERIODO) reemplaza lo que el
    negocio tenia cargado para ese mes."""
    if negocio not in NEGOCIOS:
        raise ArchivoInvalido(f"Negocio inválido: {negocio}")
    if not str(nombre_archivo or "").lower().endswith(".xlsx"):
        raise ArchivoInvalido("El archivo debe ser un Excel .xlsx")
    if len(contenido) > MAX_BYTES:
        raise ArchivoInvalido("El archivo supera los 10 MB")

    filas, avisos = _leer_archivo(negocio, contenido)
    _ensure_tables()
    periodos: list[dict] = []

    with get_connection() as cn:
        cur = cn.cursor()
        for mes in sorted({f["periodo"] for f in filas}):
            del_mes = [f for f in filas if f["periodo"] == mes]
            fecha = f"{mes}-01"
            cur.execute(
                f"""
                DELETE FROM {CIERRE}
                WHERE id_carga IN (SELECT id_carga FROM {CARGAS} WHERE negocio = ? AND periodo = ?)
                """,
                (negocio, fecha),
            )
            cur.execute(f"DELETE FROM {CARGAS} WHERE negocio = ? AND periodo = ?", (negocio, fecha))
            reemplazo = cur.rowcount > 0
            cur.execute(
                f"""
                INSERT INTO {CARGAS} (negocio, tipo, periodo, nombre_archivo, filas, cargado_por)
                OUTPUT INSERTED.id_carga
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (negocio, TIPO_CARGA, fecha, str(nombre_archivo)[:260], len(del_mes), usuario),
            )
            id_carga = cur.fetchone()[0]
            cur.executemany(
                f"""
                INSERT INTO {CIERRE}
                    (id_carga, negocio, tipo, periodo, producto, operacion, ejecutivo, ejecutivo_origen)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        id_carga,
                        negocio,
                        TIPO_CARGA,
                        fecha,
                        f["producto"][:40],
                        f["operacion"][:40],
                        f["ejecutivo"][:120],
                        f["ejecutivo_origen"][:120] or None,
                    )
                    for f in del_mes
                ],
            )
            periodos.append({"periodo": mes, "filas": len(del_mes), "reemplazo": reemplazo})
        cn.commit()
    cache.clear()

    por_ejecutivo = Counter(f["ejecutivo"] for f in filas)
    return {
        "negocio": negocio,
        "nombre_archivo": nombre_archivo,
        "filas": len(filas),
        "periodos": periodos,
        "por_producto": _por_producto(Counter(f["producto"] for f in filas)),
        "por_ejecutivo": [{"ejecutivo": e, "total": n} for e, n in sorted(por_ejecutivo.items(), key=lambda x: (-x[1], x[0]))],
        **avisos,
    }


def eliminar_carga(id_carga: int) -> None:
    _ensure_tables()
    with get_connection() as cn:
        cur = cn.cursor()
        cur.execute(f"DELETE FROM {CIERRE} WHERE id_carga = ?", (id_carga,))
        cur.execute(f"DELETE FROM {CARGAS} WHERE id_carga = ?", (id_carga,))
        cn.commit()
    cache.clear()


def _orden_producto(producto: str) -> tuple[int, str]:
    return (PRODUCTO_ORDEN.index(producto) if producto in PRODUCTO_ORDEN else len(PRODUCTO_ORDEN), producto)


def _por_producto(conteo: Counter) -> list[dict]:
    return [
        {"codigo": p, "nombre": PRODUCTO_NOMBRE.get(p, p.capitalize()), "total": conteo[p]}
        for p in sorted(conteo, key=_orden_producto)
    ]


def _conteo_mes(negocio: str, mes: str) -> list[dict]:
    _ensure_tables()
    return run_query(
        f"""
        SELECT ejecutivo, producto, COUNT_BIG(1) AS total
        FROM {CIERRE}
        WHERE negocio = ? AND periodo = CAST(? AS date)
        GROUP BY ejecutivo, producto
        """,
        (negocio, f"{mes}-01"),
    )


def get_conteo(negocio: str, periodo: str | None, ejecutivo: str | None = None) -> dict:
    """Negocios del mes por ejecutivo y producto, para la pestaña Negocios de la productividad."""
    mes = _mes(periodo)
    rows = cache.cached(("negocios", "conteo", negocio, mes), lambda: _conteo_mes(negocio, mes))
    filtro = _norm(ejecutivo)

    por_ejecutivo: dict[str, Counter] = {}
    totales: Counter = Counter()
    for row in rows:
        if filtro and _norm(row["ejecutivo"]) != filtro:
            continue
        cantidad = int(row["total"] or 0)
        por_ejecutivo.setdefault(row["ejecutivo"], Counter())[row["producto"]] += cantidad
        totales[row["producto"]] += cantidad

    filas = [
        {"ejecutivo": nombre, "conteos": dict(conteos), "total": sum(conteos.values())}
        for nombre, conteos in por_ejecutivo.items()
    ]
    filas.sort(key=lambda f: (-f["total"], f["ejecutivo"]))
    return {
        "periodo": mes,
        "productos": [{"codigo": p["codigo"], "nombre": p["nombre"]} for p in _por_producto(totales)],
        "filas": filas,
        "totales": dict(totales),
        "total": sum(totales.values()),
    }
